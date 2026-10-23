"""Tests for numbers_tool.py, run with: .venv/Scripts/python py/test_numbers_tool.py"""

import datetime as dt
import html
import os
import re
import sys
import tempfile
import zipfile
from types import SimpleNamespace

from numbers_parser import Document
from numbers_parser.constants import DurationStyle, DurationUnits, FormatType
from numbers_parser.formula import FUNCTION_NAME_TO_ID
from numbers_parser.generated import TSCEArchives_pb2 as TSCE
from numbers_parser.generated import TSDArchives_pb2 as TSD
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
import categories  # noqa: E402
import charts  # noqa: E402
import drawings  # noqa: E402
import highlights  # noqa: E402
import layout  # noqa: E402
import numbers_tool as nt  # noqa: E402
import pivot_tables  # noqa: E402
import pivots  # noqa: E402

FIXTURES = os.path.join(os.path.dirname(__file__), "..", "test", "fixtures")
OUT = tempfile.mkdtemp(prefix="numbers-tool-")


def fixture(name):
    return os.path.join(FIXTURES, name + ".numbers")


def convert(name, options=None):
    out = os.path.join(OUT, name + (f" ({options['default']})" if options else "") + ".xlsx")
    result = nt.convert(fixture(name), out, layout_options=options)
    return result, load_workbook(out)


def placed(name, sheet, options=None, table=0):
    """Where a table sits on a converted worksheet: its first, n-th, or the one with this name."""
    ws = next(ws for ws in nt.load(fixture(name)).layout(options).sheets if ws.name == sheet)
    return ws.tables[table] if isinstance(table, int) else next(p for p in ws.tables if p.table.name == table)


def at(name, ws, ref, options=None, table=0):
    """The cell at table position `ref` (e.g. "A7") on a converted sheet, allowing for the title row
    above tables whose name Numbers shows, and for where the table sits on the sheet."""
    return ws[shifted(name, ws.title, ref, options, table)]


def shifted(name, sheet, ref, options=None, table=0):
    p = placed(name, sheet, options, table)
    return nt.shift_rows(ref, p.row, p.col)


def grid_cell(sheet, p, r, c):
    """A table's cell (r, c) in a previewed worksheet's grid."""
    return sheet["rows"][p.row + r][p.col + c]


# Formulas naming a table "Prices" on sheet "Main", and a table "Totals" alone on its sheet.
FakeLayout = SimpleNamespace(
    by_table={("Main", "Prices"): 1, (None, "Prices"): 1, (None, "Totals"): 2}, by_sheet={"Totals": 2},
    place={1: layout.Placed(None, "Main - Prices", 0, 0, None), 2: layout.Placed(None, "Totals", 0, 0, None)})


def test_date_patterns():
    cases = {
        "MMM d, y": "mmm d, yyyy",
        "EEEE, MMMM d, y h:mm a": "dddd, mmmm d, yyyy h:mm AM/PM",
        "yyyy-MM-dd": "yyyy-mm-dd",
        "HH:mm:ss.SSS": "hh:mm:ss.000",
        "d 'de' MMMM": 'd "de" mmmm',
        "'Day' D 'of' yyyy": "m/d/yyyy",  # day-of-year has no Excel equivalent
        "K:mm": "m/d/yyyy h:mm",
        "hhmmssSSS a": "m/d/yyyy h:mm",  # Excel only accepts fractional seconds after "s."
        "": "m/d/yyyy",
    }
    for pattern, expected in cases.items():
        assert nt.icu_to_excel_date(pattern) == expected, (pattern, nt.icu_to_excel_date(pattern))


def test_formula_translation():
    t = lambda f: nt.translate_formula(f, FakeLayout, "Main")  # noqa: E731
    assert t("SUM(A1:B2)") == "=SUM(A1:B2)"
    assert t("A1÷B1×2") == "=A1/B1*2"
    assert t("IF(A1≥2,\"big\",\"small\")") == '=IF(A1>=2,"big","small")'
    assert t("SUM(Prices::B2:B9)") == "=SUM('Main - Prices'!B2:B9)"
    assert t("Main::Prices::C3*2") == "='Main - Prices'!C3*2"
    assert t("SUM((A1+B1),C1)") == "=SUM((A1+B1),C1)"
    assert t("TODAY()") == "=TODAY()"
    assert t("IF(A1>0,1)") == "=IF(A1>0,1)"
    assert t('"Row "&CHAR(64+COLUMN())&ROW()') == '="Row "&CHAR(64+COLUMN())&ROW()'
    # Things Excel would reject or compute differently keep their values instead.
    assert t("IF(A1)") is None  # too few arguments
    assert t("PI(1)") is None  # too many arguments
    assert t("DURATION(1,2,3)") is None  # same name, different meaning in Excel
    assert t("SUM(A:A)") is None  # whole-column reference
    assert t("AVERAGEIFS({1,2},{1,0},\">0\")") is None  # array constant
    assert t("COUNTA(fruit)") is None  # header name nothing defines
    assert t("SUM(Unknown::A1)") is None
    assert t("STOCK(\"AAPL\",1)") is None
    assert t("GCD(5,(A1,A2))") is None  # list syntax
    assert t("DATEDIF(\"15/05/2001\",\"15/07/2002\",\"D\")") is None  # region-dependent text dates


def test_header_names_and_titles_in_formulas():
    # In Numbers a header name means the column's body cells inside a function call
    # (COUNTA('Column X & Y') counts 7 cells) and the same row's cell elsewhere. The table also
    # moves down a row in Excel under its title, which references and ROW() must allow for.
    book = nt.load(fixture("test-styles"))
    sheet_name, table = book.tables[1]
    lay = book.layout()
    assert lay.place[table._table_id].row == 1
    cell = next(c for row in nt.table_rows(table) for c in row if c.is_formula)
    assert cell.formula == '"Row Header "&CHAR(64+COLUMN())&ROW()'
    assert nt.translate_formula(cell.formula, lay, sheet_name, cell) == '="Row Header "&CHAR(64+COLUMN())&(ROW()-1)'

    _, wb = convert("create-formulas")
    ws = wb["Main Sheet - Formula Tests"]
    assert at("create-formulas", ws, "A7").value == "=SUM(" + shifted("create-formulas", ws.title, "A2:B2") + ")"


def test_approximate_lookups_need_sorted_data():
    # Excel's approximate MATCH/LOOKUP/VLOOKUP assume sorted data; Numbers doesn't. Unsorted -> value.
    book = SimpleNamespace(table_by_id={})
    rows = {"sorted": [1, 2, 5, 9, None], "unsorted": ["Pears", "Apples", "Leeks"]}

    def check(name, mode, values):
        table = SimpleNamespace(num_rows=len(values) + 1, num_cols=1, _table_id=7,
                                _model=SimpleNamespace(objects={7: SimpleNamespace(
                                    number_of_header_rows=1, number_of_footer_rows=0, number_of_header_columns=0,
                                    HasField=lambda f: False)}))
        book.table_by_id[7] = table
        original = nt.table_rows
        nt.table_rows = lambda t: [[SimpleNamespace(value="Header")]] + [[SimpleNamespace(value=v)] for v in values]
        try:
            args = [[("str", '"x"')], [("exact", "⟦7|||0||0⟧")]] + ([[mode]] if mode else [])  # all of column A
            return nt.lookup_is_safe(name, args, book)
        finally:
            nt.table_rows = original

    assert check("MATCH", None, rows["sorted"])  # approximate on sorted data, trailing blank ignored
    assert not check("MATCH", None, rows["unsorted"])
    assert check("MATCH", ("num", "0"), rows["unsorted"])  # exact match is always fine
    assert not check("LOOKUP", None, rows["unsorted"])


def test_styles_are_kept():
    result, wb = convert("test-styles")
    assert result["sheets"] == 4
    assert result["formulas"] == 41 and result["formulasAsValues"] == 0
    ws = wb["Styles"]
    title = at("test-styles", ws, "A1")
    assert title.value == 'name="Title",bold,font_size=24.0'
    assert title.font.b and title.font.sz == 24
    red = next(c for row in ws.iter_rows() for c in row if c.value and "Heading Red" in str(c.value))
    assert red.font.color.rgb.endswith("EE220C")
    strike = next(c for row in ws.iter_rows() for c in row if c.value and "strikethrough" in str(c.value) and "font" not in str(c.value))
    assert strike.font.strike


def test_fills_and_borders_are_kept():
    _, wb = convert("test-bgcolour")
    fills = {c.fill.fgColor.rgb for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.fill.fill_type == "solid"}
    assert len(fills) > 1
    _, wb = convert("test-border-widths")
    borders = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.border.top.style or c.border.left.style]
    assert borders


def test_formulas_are_kept():
    result, wb = convert("create-formulas")
    assert result["formulas"] > 30
    ws = wb["Main Sheet - Formula Tests"]
    assert at("create-formulas", ws, "A7").value == "=SUM(" + shifted("create-formulas", ws.title, "A2:B2") + ")"


def test_preview():
    data = nt.preview(fixture("test-styles"))
    sheet = data["sheets"][0]
    assert sheet["name"] == "Styles"
    assert sheet["colWidths"] and sheet["rowHeights"]
    first = grid_cell(sheet, placed("test-styles", "Styles"), 0, 0)
    text, style = first
    assert text.startswith('name="Title"')
    assert sheet["styles"][style]["b"] == 1 and sheet["styles"][style]["size"] == 24


def test_theme_styles_follow_table_regions():
    # Cells styled only by the table theme must take the header/body/header-column style of
    # where they are, not whichever theme-styled cell happened to be read first.
    sheet = nt.preview(fixture("test-pivot"))["sheets"][0]
    p = placed("test-pivot", sheet["name"])
    cell = lambda r, c: grid_cell(sheet, p, r, c)  # noqa: E731
    style = lambda r, c: sheet["styles"][cell(r, c)[1]] if isinstance(cell(r, c), list) else {}  # noqa: E731
    assert style(0, 1).get("bg") == "#BEC0BF"  # header row
    assert style(1, 0).get("bg") == "#DCDCDC"  # header column
    assert "bg" not in style(1, 1)  # body


def test_text_is_never_invisible():
    # Light text only on a dark background, in every fixture.
    for name in [f[:-8] for f in os.listdir(FIXTURES) if f.endswith(".numbers") and f != "encrypted.numbers"]:
        for sheet in nt.preview(fixture(name))["sheets"]:
            for s in sheet["styles"]:
                if "color" in s and nt.is_light(s["color"]):
                    assert not nt.is_light(s.get("bg", "#FFFFFF")), (name, sheet["name"], s)


def test_pivot_values_match_numbers():
    # This file was saved by a Numbers version that also stored the pivot results, so the
    # recalculation can be checked against Numbers' own numbers.
    doc = Document(fixture("test-pivot"))
    checked = 0
    for sheet in doc.sheets:
        for table in sheet.tables:
            pivot = pivots.pivot_for(table, recompute_stored=True)
            if pivot is None:
                continue
            assert pivot.missing == 0
            for (r, c), (value, _) in pivot.values.items():
                assert abs(table.cell(r, c).value - value) < 1e-6, (sheet.name, r, c)
                checked += 1
    assert checked == 28


def pivot_parts(path):
    """{worksheet name: [(pivotTable xml, its pivotCacheDefinition xml, its pivotCacheRecords xml)]}."""
    found = {}
    with zipfile.ZipFile(path) as z:
        read = lambda name: z.read(name).decode()  # noqa: E731
        targets = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', read("xl/_rels/workbook.xml.rels")))
        for name, rid in re.findall(r'<sheet name="([^"]+)" sheetId="\d+" r:id="(rId\d+)"', read("xl/workbook.xml")):
            rels = "xl/" + targets[rid].replace("worksheets/", "worksheets/_rels/") + ".rels"
            if rels not in z.namelist():
                continue
            for table in re.findall(r'Target="\.\./pivotTables/([^"]+)"', read(rels)):
                cache = re.search(r'Target="\.\./pivotCache/([^"]+)"', read(f"xl/pivotTables/_rels/{table}.rels")).group(1)
                found.setdefault(html.unescape(name), []).append(
                    (read(f"xl/pivotTables/{table}"), read(f"xl/pivotCache/{cache}"),
                     read("xl/pivotCache/" + cache.replace("Definition", "Records"))))
    return found


def grid_values(ws, p, rows, cols):
    """A placed table's cells on a converted worksheet, as a list of rows of values."""
    return [[ws.cell(p.row + r + 1, p.col + c + 1).value for c in range(cols)] for r in range(rows)]


def test_pivot_tables_convert():
    # Numbers' pivot tables become Excel PivotTables reading the converted source table, laid out
    # as Excel lays them out, with the totals Numbers shows (its own picture of this file shows
    # APAC 84,800, EMEA 83,700, NA 95,400 and 263,900 in all) before anyone refreshes them.
    result, wb = convert("test-pivot")
    assert result["unsupported"]["pivots"] == 0 and result["pivots"] == 2
    ws = wb["Sales - Sales Pivot"]
    p = placed("test-pivot", ws.title)
    assert p.size == (6, 6)  # Excel's layout: two header rows, three regions and a grand total row
    assert grid_values(ws, p, 6, 6) == [
        ["Revenue (Sum)", "Quarter", None, None, None, None],
        ["Region", "Q1", "Q2", "Q3", "Q4", "Grand Total"],
        ["APAC", 17700, 20400, 28900, 17800, 84800],
        ["EMEA", 18400, 22900, 23800, 18600, 83700],
        ["NA", 12400, 27400, 28600, 27000, 95400],
        ["Grand Total", 48500, 70700, 81300, 63400, 263900]]
    assert ws.cell(p.row + 4, p.col + 2).number_format == "#,##0"  # as Numbers shows the pivot's values
    formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=" not in formulas  # undecodable pivot formulas must not become empty formulas

    (table, cache, records), = pivot_parts(os.path.join(OUT, "test-pivot.xlsx"))["Sales - Sales Pivot"]
    source = placed("test-pivot", "Sales - Sales")
    ref = nt.shift_rows("A1:C13", source.row, source.col)  # the header row and the twelve rows, not the total
    assert f'<worksheetSource ref="{ref}" sheet="Sales - Sales"/>' in cache and 'recordCount="12"' in cache
    assert '<cacheField name="Region" numFmtId="0"><sharedItems count="3"><s v="EMEA"/><s v="APAC"/><s v="NA"/>' in cache
    assert records.count("<r>") == 12 and '<r><x v="0"/><x v="0"/><n v="18400"/></r>' in records
    assert f'<location ref="A{p.row + 1}:F{p.row + 6}" firstHeaderRow="1" firstDataRow="2" firstDataCol="1"/>' in table
    assert '<pivotField axis="axisRow" compact="0" outline="0" showAll="0" sortType="ascending"><items count="4">' \
           '<item x="1"/><item x="0"/><item x="2"/><item t="default"/></items></pivotField>' in table  # APAC, EMEA, NA
    assert '<rowFields count="1"><field x="0"/></rowFields>' in table and '<colFields count="1"><field x="1"/></colFields>' in table
    fmt = re.search(r'<dataField name="Revenue \(Sum\)" fld="2" baseField="0" baseItem="0" numFmtId="(\d+)"/>', table).group(1)
    assert "rowGrandTotals" not in table and "colGrandTotals" not in table  # Numbers shows both
    assert 'pivotTableStyleInfo name="Numbers Pivot 1"' in table
    styles = zipfile.ZipFile(os.path.join(OUT, "test-pivot.xlsx")).read("xl/styles.xml").decode()
    assert f'<numFmt numFmtId="{fmt}" formatCode="#,##0"/>' in styles  # the values' number format, kept on refresh
    assert '<tableStyle name="Numbers Pivot 1" table="0"' in styles and 'pivotButton="1"' in styles
    header = re.search(r'type="headerRow" dxfId="(\d+)"', styles).group(1)
    assert 'rgb="FFBEC0BF"' in re.findall(r"<dxf>.*?</dxf>", styles)[int(header)]  # Numbers' header row colour
    pivots_read = {ws.title: ws._pivots for ws in wb.worksheets if ws._pivots}  # openpyxl reads them back
    assert [p.name for p in pivots_read["Sales - Sales Pivot"]] == ["Sales Pivot"]


def test_pivot_tables_on_one_sheet():
    # Laid out as in Numbers, the PivotTable shares a worksheet with its source table and takes the
    # room its totals need, so the table below it moves down rather than being covered.
    options = {"default": "sheet"}
    result, wb = convert("test-pivot", options)
    assert result["pivots"] == 2
    ws = wb["Sales"]
    pivot, source = placed("test-pivot", "Sales", options, "Sales Pivot"), placed("test-pivot", "Sales", options, "Sales")
    metrics = placed("test-pivot", "Sales", options, "Metrics")
    assert pivot.size == (6, 6) and metrics.row > pivot.row + 6 and metrics.col == pivot.col
    assert grid_values(ws, pivot, 6, 6)[5] == ["Grand Total", 48500, 70700, 81300, 63400, 263900]
    (table, cache, _), = pivot_parts(os.path.join(OUT, "test-pivot (sheet).xlsx"))["Sales"]
    ref = nt.shift_rows("A1:C13", source.row, source.col)
    assert f'<worksheetSource ref="{ref}" sheet="Sales"/>' in cache
    assert f'<location ref="{nt.shift_rows("A1:F6", pivot.row, pivot.col)}"' in table


def test_pivot_preview():
    # The preview shows the PivotTable as Excel will, with drop-down arrows on its field buttons.
    sheet = next(s for s in nt.preview(fixture("test-pivot"))["sheets"] if s["name"] == "Sales - Sales Pivot")
    p = placed("test-pivot", sheet["name"])
    cell = lambda r, c: grid_cell(sheet, p, r, c)  # noqa: E731
    text = lambda r, c: cell(r, c) if isinstance(cell(r, c), str) else cell(r, c)[0]  # noqa: E731
    assert [text(1, c) for c in range(6)] == ["Region", "Q1", "Q2", "Q3", "Q4", "Grand Total"]
    assert [text(r, 5) for r in range(2, 6)] == ["84,800", "83,700", "95,400", "263,900"]
    assert cell(1, 0)[2] == {"control": "popup"} and cell(0, 1)[2] == {"control": "popup"}  # Region and Quarter
    assert sheet["styles"][cell(1, 1)[1]]["bg"] == "#BEC0BF" and sheet["styles"][cell(2, 0)[1]]["bg"] == "#DCDCDC"
    assert sheet["styles"][cell(2, 5)[1]]["b"] == 1  # the grand total column, bold as in Numbers


def view_rows(table, order):
    """Makes Numbers show a pivot table's rows in another order: `order` lists the rows' indexes in
    the table, "grand" for the grand total row (left out to hide it)."""
    model = table._model
    view = model.objects[pivots.find_table_info(model, table._table_id).view_column_row_uids.identifier]
    base = model.objects[model.objects[table._table_id].base_column_row_uids.identifier]
    uids = {base.row_index_for_uid[i]: (u.lower, u.upper) for i, u in enumerate(base.sorted_row_uids)}
    uids["grand"] = pivot_tables.GRAND_UID
    del view.sorted_row_uids[:], view.row_index_for_uid[:], view.row_uid_for_index[:]
    for position, row in enumerate(order):
        uid = view.sorted_row_uids.add()
        uid.lower, uid.upper = uids[row]
        view.row_index_for_uid.append(position)
        view.row_uid_for_index.append(position)


def test_pivot_order_and_totals_follow_numbers():
    # The order Numbers shows a pivot's groups in, and whether it shows a grand total, come from
    # what it last showed. (In this file the rows are EMEA, APAC, NA in the table and Numbers sorts
    # them A to Z; here it's made to show them Z to A, then in no order, without a grand total.)
    book = nt.load(fixture("test-pivot"))
    _, table = book.tables[1]
    try:
        for order, sort, items in (([0, 1, 4, 2, 3], "descending", [2, 0, 1]), ([0, 1, 3, 4, 2], None, [1, 2, 0])):
            view_rows(table, order)
            nt.native_pivot(book, table, pivots.pivot_for(table))
            book.layouts.clear()
            out = os.path.join(OUT, "pivot-order.xlsx")
            nt.convert(fixture("test-pivot"), out)
            (xml, _, _), = pivot_parts(out)["Sales - Sales Pivot"]
            field = re.search(r'<pivotField axis="axisRow"[^>]*>.*?</pivotField>', xml).group(0)
            assert (f'sortType="{sort}"' in field) if sort else "sortType" not in field
            assert re.findall(r'<item x="(\d)"/>', field) == [str(x) for x in items]
            assert 'rowGrandTotals="0"' in xml and "colGrandTotals" not in xml
            ws = load_workbook(out)["Sales - Sales Pivot"]
            p = book.layout().place[table._table_id]
            names = ["EMEA", "APAC", "NA"]
            assert [ws.cell(p.row + 3 + i, p.col + 1).value for i in range(3)] == [names[x] for x in items]
            assert ws.cell(p.row + 6, p.col + 1).value is None and p.size == (5, 6)  # no grand total row
        # Groups sorted by their totals (NA 95,400, APAC 84,800, EMEA 83,700) use Excel's sort by
        # value, which keeps them in that order when the data changes.
        for order, sort, shown in (([0, 1, 4, 3, 2, "grand"], "descending", ["NA", "APAC", "EMEA"]),
                                   ([0, 1, 2, 3, 4, "grand"], "ascending", ["EMEA", "APAC", "NA"])):
            view_rows(table, order)
            nt.native_pivot(book, table, pivots.pivot_for(table))
            book.layouts.clear()
            out = os.path.join(OUT, "pivot-by-value.xlsx")
            nt.convert(fixture("test-pivot"), out)
            (xml, _, _), = pivot_parts(out)["Sales - Sales Pivot"]
            field = re.search(r'<pivotField axis="axisRow"[^>]*>.*?</pivotField>', xml).group(0)
            assert f'sortType="{sort}"' in field and '<reference field="4294967294" count="1" selected="0"><x v="0"/>' in field
            ws = load_workbook(out)["Sales - Sales Pivot"]
            p = book.layout().place[table._table_id]
            assert [ws.cell(p.row + 3 + i, p.col + 1).value for i in range(3)] == shown
    finally:
        nt._cache.clear()


def test_pivot_summaries():
    # How a value is summarised is read from its label ("Revenue (Average)"); Excel's PivotTable
    # summarises it the same way.
    book = nt.load(fixture("test-pivot"))
    _, table = book.tables[1]
    try:
        owner = table._model.objects[table._model.objects[table._table_id].pivot_owner.identifier]
        owner.aggregate_columns.aggregates[0].agg_type = 99  # a code that doesn't say
        table.cell(1, 1)._value = "Revenue (Average)"
        nt.native_pivot(book, table, pivots.pivot_for(table))
        book.layouts.clear()
        out = os.path.join(OUT, "pivot-average.xlsx")
        nt.convert(fixture("test-pivot"), out)
        (xml, _, _), = pivot_parts(out)["Sales - Sales Pivot"]
        assert '<dataField name="Revenue (Average)" fld="2" subtotal="average"' in xml
        ws = load_workbook(out)["Sales - Sales Pivot"]
        p = book.layout().place[table._table_id]
        assert ws.cell(p.row + 3, p.col + 6).value == 84800 / 4 and abs(ws.cell(p.row + 6, p.col + 6).value - 263900 / 12) < 1e-9
        assert ws.cell(p.row + 3, p.col + 6).number_format == "#,##0.00"  # two decimal places for averages of whole numbers
        table.cell(1, 1)._value = "Revenue (Running Total)"  # a summary Excel can't match: values, as before
        nt.native_pivot(book, table, pivots.pivot_for(table))
        assert table._table_id not in book.native_pivots
    finally:
        nt._cache.clear()


def test_pivot_date_groupings():
    # A Numbers grouping that splits dates exactly as one of Excel's does becomes that grouping
    # (Excel then groups new dates itself); any other is listed group by group, named as in Numbers.
    month = SimpleNamespace(grouping_functor=SimpleNamespace(formula=SimpleNamespace(AST_node_array=SimpleNamespace(
        AST_node=rpn("x", "SUM/1", "x", "=", "x", "ISBLANK/1", "NOT/1", "AND/2",
                     "{", 2020, "x", "MONTH/1", 1, "DATE/3", "}", "{", 1, 0, "/", "}", "IF/3")))))
    dates = [dt.datetime(2021, 1, 5), dt.datetime(2021, 3, 1), None, dt.datetime(2022, 1, 9)]
    keys = {0: pivots.norm(dt.datetime(2020, 1, 1)), 1: pivots.norm(dt.datetime(2020, 3, 1)), 2: pivot_tables.Excluded, 3: None}
    keys[3] = keys[0]
    group = pivot_tables.native_group(month, dates, keys)
    assert group.by == "months" and group.items[1:13] == pivot_tables.MONTHS and group.items[0] == "<1/5/2021"
    assert group.mapping == [1, 3, 0, 1] and group.hidden == {0}  # Numbers leaves blank dates out; Excel puts them first
    weekday = SimpleNamespace(grouping_functor=SimpleNamespace(formula=SimpleNamespace(AST_node_array=SimpleNamespace(
        AST_node=rpn("x", "WEEKDAY/1")))))
    assert pivot_tables.native_group(weekday, dates, keys) is None
    listed = pivot_tables.discrete_group(dates, {0: 3, 1: 2, 2: pivot_tables.Excluded, 3: 3}, {3: "Tuesday", 2: "Monday"})
    assert listed.items == ["Tuesday", "Monday", "(blank)"] and listed.mapping == [0, 1, 2, 0] and listed.hidden == {2}


def synthetic_pivot(columns, records, rows, cols, data, data_on_rows=False, no_subtotals=()):
    """A pivot_tables.Spec over `records` (one value per column) with the row and column fields
    given by column index (items in Excel's A to Z order) and values as (column, kind, name),
    laid out."""
    spec = pivot_tables.Spec()
    spec.records, spec.included, spec.data_on_rows = len(records), [True] * len(records), data_on_rows
    spec.fields = [pivot_tables.CacheField(name, [r[i] for r in records]) for i, name in enumerate(columns)]

    def axis(i):
        shared = spec.fields[i].share()
        order = sorted(range(len(shared)), key=lambda x: pivot_tables.sort_key(shared[x]))
        position = {x: p for p, x in enumerate(order)}
        made = pivot_tables.Axis(i, [{"x": x, "caption": None, "hidden": False, "key": None} for x in order],
                                 [position[x] for x in spec.fields[i].index])
        made.subtotals = i not in no_subtotals
        return made
    spec.rows, spec.cols = [axis(i) for i in rows], [axis(i) for i in cols]
    spec.data = [pivot_tables.DataField(i, kind, name, None) for i, kind, name in data]
    spec.values = [spec.fields[d.field].values for d in spec.data]
    pivot_tables.render(spec)
    return spec


def test_pivot_layout_matches_excel():
    # Excel checks a PivotTable's list of row and column items against its layout when it opens
    # the file. These are the lists, layouts and cells Excel itself saved for the same pivots.
    rows = [("a1", "b1", "c1", "d1", 1.0), ("a1", "b1", "c2", "d2", 2.0), ("a1", "b2", "c1", "d1", 3.0),
            ("a2", "b1", "c1", "d2", 4.0), ("a2", "b2", "c2", "d1", 5.0), ("a2", "b2", "c1", "d1", 6.0),
            ("a1", "b2", "c2", "d2", 7.0)]
    names = ("A", "B", "C", "D", "V")
    spec = synthetic_pivot(names, rows, [0, 1, 2], [3], [(4, "sum", "V (Sum)")])
    g = spec.grid
    assert (g.rows, g.cols, g.first_header_row, g.first_data_row, g.first_data_col) == (16, 6, 1, 2, 3)
    assert pivot_tables.items_xml(g.row_items) == (
        'count="14"><i><x/><x/><x/></i><i r="2"><x v="1"/></i><i t="default" r="1"><x/></i><i r="1"><x v="1"/><x/></i>'
        '<i r="2"><x v="1"/></i><i t="default" r="1"><x v="1"/></i><i t="default"><x/></i><i><x v="1"/><x/><x/></i>'
        '<i t="default" r="1"><x/></i><i r="1"><x v="1"/><x/></i><i r="2"><x v="1"/></i><i t="default" r="1"><x v="1"/></i>'
        '<i t="default"><x v="1"/></i><i t="grand"><x/></i>')
    text = lambda r, c: g.cells[(r, c)].value if (r, c) in g.cells else None  # noqa: E731
    assert [text(0, c) for c in range(6)] == ["V (Sum)", None, None, "D", None, None]
    assert [text(4, c) for c in range(6)] == [None, "b1 Total", None, 1, 2, 3]
    assert [text(15, c) for c in range(6)] == ["Grand Total", None, None, 15, 13, 28]

    spec = synthetic_pivot(names, rows, [0, 1], [], [(4, "sum", "V (Sum)"), (4, "average", "V (Average)")], data_on_rows=True)
    g = spec.grid
    assert (g.rows, g.cols, g.first_data_row, g.first_data_col) == (15, 4, 1, 3)
    assert pivot_tables.items_xml(g.row_items).startswith(
        'count="14"><i><x/><x/><x/></i><i r="2" i="1"><x v="1"/></i><i r="1"><x v="1"/><x/></i><i r="2" i="1"><x v="1"/></i>'
        '<i t="default"><x/></i><i t="default" i="1"><x/></i>')
    assert [g.cells[(5, 0)].value, g.cells[(14, 0)].value, g.cells[(14, 3)].value] == ["a1 V (Sum)", "Total V (Average)", 4]

    spec = synthetic_pivot(names, rows, [], [0, 3], [(4, "sum", "V (Sum)"), (4, "count", "V (Count)")])
    g = spec.grid
    assert (g.rows, g.cols, g.first_header_row, g.first_data_row, g.first_data_col) == (5, 14, 1, 4, 0)
    group = '<i r="1"><x v="1"/><x/></i><i r="2" i="1"><x v="1"/></i>'
    assert pivot_tables.items_xml(g.col_items) == (
        'count="14"><i><x/><x/><x/></i><i r="2" i="1"><x v="1"/></i>' + group + '<i t="default"><x/></i><i t="default" i="1"><x/></i>'
        '<i><x v="1"/><x/><x/></i><i r="2" i="1"><x v="1"/></i>' + group + '<i t="default"><x v="1"/></i>'
        '<i t="default" i="1"><x v="1"/></i><i t="grand"><x/></i><i t="grand" i="1"><x/></i>')
    assert pivot_tables.items_xml(g.row_items) == 'count="1"><i/>'
    assert [g.cells[(1, c)].value for c in (4, 5, 12, 13)] == ["a1 V (Sum)", "a1 V (Count)", "Total V (Sum)", "Total V (Count)"]
    assert [g.cells[(4, c)].value for c in (4, 5, 12, 13)] == [13, 4, 28, 7]

    spec = synthetic_pivot(names, rows, [0, 1, 2], [], [(4, "sum", "V (Sum)")], no_subtotals=(1,))
    assert [e[0] for e in spec.grid.row_items].count("default") == 2  # only A's groups have subtotals


def test_pivot_cache_items():
    # The pivot cache says what each field holds as Excel does, or Excel asks to repair the file.
    shared = pivot_tables.shared_items_xml
    assert shared([1.0, "text", None, 2.5, 3.0, "zz"], None) == \
        '<sharedItems containsBlank="1" containsMixedTypes="1" containsNumber="1" minValue="1" maxValue="3"/>'
    assert shared([211.0, 476.0], None) == \
        '<sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1" containsInteger="1" minValue="211" maxValue="476"/>'
    when = [dt.datetime(2021, 1, 1, 6), None, dt.datetime(2021, 1, 6, 12)]
    assert shared(when, when) == ('<sharedItems containsNonDate="0" containsDate="1" containsString="0" containsBlank="1" '
                                  'minDate="2021-01-01T06:00:00" maxDate="2021-01-06T12:00:00" count="3">'
                                  '<d v="2021-01-01T06:00:00"/><m/><d v="2021-01-06T12:00:00"/></sharedItems>')
    assert shared(["A", "B"], ["A", "B"]) == '<sharedItems count="2"><s v="A"/><s v="B"/></sharedItems>'
    assert pivot_tables.value_xml("a\nb & <c>") == '<s v="a&#10;b &amp; &lt;c&gt;"/>'


def rpn(*tokens):
    """Build a grouping-formula node list: numbers, strings, 'x', 'FUNC/n', operators, '{' and '}' for thunks."""
    nodes = []
    for tok in tokens:
        n = TSCE.ASTNodeArrayArchive.ASTNodeArchive()
        if isinstance(tok, (int, float)):
            n.AST_node_type, n.AST_number_node_number = pivots.NODE.Value("NUMBER_NODE"), tok
        elif tok == "x":
            n.AST_node_type, n.AST_symbol = pivots.NODE.Value("VAR_NODE"), 0
        elif tok in ("{", "}"):
            n.AST_node_type = pivots.NODE.Value("BEGIN_THUNK_NODE" if tok == "{" else "END_THUNK_NODE")
        elif "/" in tok and tok.split("/")[0].isupper():
            name, argc = tok.split("/")
            n.AST_node_type = pivots.NODE.Value("FUNCTION_NODE")
            n.AST_function_node_index = FUNCTION_NAME_TO_ID[name]
            n.AST_function_node_numArgs = int(argc)
        elif tok in ("+", "-", "*", "/", "=", ">", "<"):
            names = {"+": "ADDITION", "-": "SUBTRACTION", "*": "MULTIPLICATION", "/": "DIVISION", "=": "EQUAL_TO",
                     ">": "GREATER_THAN", "<": "LESS_THAN"}
            n.AST_node_type = pivots.NODE.Value(names[tok] + "_NODE")
        else:
            n.AST_node_type, n.AST_string_node_string = pivots.NODE.Value("STRING_NODE"), tok
        nodes.append(n)
    return nodes


def test_pivot_grouping_formulas():
    d = dt.datetime(2021, 3, 17, 14, 30)
    # Numbers' own "by month" rule: IF(AND(SUM(x)=x, NOT(ISBLANK(x))), DATE(2020, MONTH(x), 1), 1/0)
    month = rpn("x", "SUM/1", "x", "=", "x", "ISBLANK/1", "NOT/1", "AND/2",
                "{", 2020, "x", "MONTH/1", 1, "DATE/3", "}", "{", 1, 0, "/", "}", "IF/3")
    assert pivots.evaluate(month, d) == dt.datetime(2020, 3, 1)
    for bad in (None, "text"):
        try:
            pivots.evaluate(month, bad)
        except pivots.GroupError:
            pass
        else:
            raise AssertionError(bad)
    quarter = rpn("Q", "x", "MONTH/1", 2, "+", 3, "/", "INT/1", "CONCATENATE/2")
    assert pivots.evaluate(quarter, d) == "Q1"
    assert pivots.evaluate(rpn("x", "YEAR/1"), d) == 2021
    assert pivots.evaluate(rpn("x", "WEEKDAY/1"), d) == 4  # Wednesday, Sunday = 1
    assert pivots.evaluate(rpn("x", 2, "WEEKDAY/2"), d) == 3  # Monday = 1
    assert pivots.evaluate(rpn("x", "WEEKNUM/1"), dt.datetime(2021, 1, 2)) == 1
    assert pivots.evaluate(rpn("x", "WEEKNUM/1"), dt.datetime(2021, 1, 3)) == 2  # first Sunday starts week 2
    assert pivots.evaluate(rpn("x", 10, "FLOOR/2"), 37) == 30
    try:
        pivots.evaluate(rpn("x", "STOCK/1"), d)
    except pivots.Unsupported:
        pass
    else:
        raise AssertionError("unknown functions must not be guessed")


def test_pivot_aggregates():
    values = [3, None, 5.5, "n/a", 1.5]
    assert pivots.aggregate("sum", values) == 10
    assert pivots.aggregate("average", values) == 10 / 3
    assert pivots.aggregate("count", values) == 4
    assert pivots.aggregate("min", values) == 1.5 and pivots.aggregate("max", values) == 5.5
    assert pivots.aggregate("sum", [None, "x"]) is None


def test_duration_formats():
    fmt = lambda style, largest, smallest, auto=False: SimpleNamespace(  # noqa: E731
        duration_style=style, duration_unit_largest=largest, duration_unit_smallest=smallest,
        use_automatic_duration_units=auto)
    S, L, C = DurationStyle.SHORT, DurationStyle.LONG, DurationStyle.COMPACT
    U = DurationUnits
    two_days = 2 * 86400 + 4 * 3600 + 20 * 60 + 2
    assert nt.duration_format(fmt(S, U.DAY, U.SECOND), two_days) == 'd"d" h"h" m"m" s"s"'
    assert nt.duration_format(fmt(C, U.HOUR, U.SECOND), two_days) == "[h]:mm:ss"
    assert nt.duration_format(fmt(L, U.HOUR, U.MINUTE), 3600 + 60) == '[h]" hour" m" minute"'
    assert nt.duration_format(fmt(S, U.DAY, U.DAY), 60 * 86400) == '0"d"'  # any number of days
    assert nt.duration_format(fmt(S, U.DAY, U.HOUR), 40 * 86400) == '[h]"h"'  # past 31 days: hours
    assert nt.duration_format(nt.DEFAULT_DURATION_FORMAT, 10 * 86400) == '0"d"'


def test_number_formats():
    fmt = lambda **kw: SimpleNamespace(**{"decimal_places": 253, "show_thousands_separator": False,  # noqa: E731
                                          "negative_style": 0, "currency_code": "USD", "use_accounting_style": False,
                                          "fraction_accuracy": 0, **kw})
    assert nt.builtin_number_format(fmt(format_type=FormatType.PERCENT), 6.43) == "0%"
    assert nt.builtin_number_format(fmt(format_type=FormatType.PERCENT), 0.0533) == "0.00%"
    assert nt.builtin_number_format(fmt(format_type=FormatType.DECIMAL, show_thousands_separator=True), 1234.56) == "#,##0.00"
    assert nt.builtin_number_format(fmt(format_type=FormatType.FRACTION, fraction_accuracy=0xFFFFFFFE), 12.34) == "# ??/??"
    assert nt.builtin_number_format(fmt(format_type=FormatType.CURRENCY, decimal_places=2, use_accounting_style=True), 5) == '"$"* #,##0.00'


def test_format_number():
    assert nt.format_number(1234.5, '"$"#,##0') == "$1,235"
    assert nt.format_number(-5, '"$"#,##0.00') == "-$5.00"
    assert nt.format_number(0.256, "0.0%") == "25.6%"
    assert nt.format_number(476.0, None) == "476"


def chart_on(ws, kind=None):
    """The charts on a converted sheet, left to right, optionally only those of one kind."""
    found = sorted(ws._charts, key=lambda c: (c.anchor._from.col, c.anchor._from.colOff))
    return [c for c in found if kind is None or c.tagname == kind]


def test_charts_convert():
    # issue-9 has a chart of each kind Numbers offers. Excel has no bubble charts, so that one is reported.
    result, wb = convert("issue-9")
    assert result["charts"] == 10 and result["unsupported"]["charts"] == 1
    # Charts read the table cells, so they stay live: here B2:B6 of the table, below its title row.
    ws = wb["Column and Bar Charts"]
    column, bar = chart_on(ws, "barChart")
    top = placed("issue-9", ws.title).row
    assert bar.barDir == "bar" and column.barDir == "col"
    assert bar.series[0].val.numRef.f == f"'Column and Bar Charts'!$B${2 + top}:$B${6 + top}"
    assert bar.series[0].cat.strRef.f == f"'Column and Bar Charts'!$A${2 + top}:$A${6 + top}"
    assert [ws[f"B{r + top}"].value for r in range(2, 7)] == [11, 15, 9, 14, 21]
    assert bar.x_axis.scaling.orientation == "maxMin"  # Numbers lists bars from the top
    # One data set, one bar per series: drawn as one series with a colour per bar, named underneath.
    assert len(column.series) == 1 and len(column.series[0].dPt) == 5
    assert column.series[0].dLbls.showVal  # Numbers shows the values on the bars

    pie, = chart_on(wb["Pie and Donut Charts"], "pieChart")
    donut, = chart_on(wb["Pie and Donut Charts"], "doughnutChart")
    assert len(pie.series) == 1 and len(pie.series[0].dPt) == 5 and pie.series[0].dLbls.showPercent
    assert donut.holeSize == 50

    combined, = chart_on(wb["2-Axis Chart"])
    assert {c.tagname for c in combined._charts} == {"barChart", "lineChart"}
    scatter, = chart_on(wb["Scatter Chart"])
    assert scatter.series[0].xVal.numRef.f.startswith("'Scatter Chart'!$A$")

    # A chart above its table gets rows of its own height above it; the table and every
    # reference to it move down to match.
    ws = wb["Interactive Chart"]
    interactive, = chart_on(ws)
    p = placed("issue-9", ws.title)
    assert p.title_row == 1 and p.row == 2  # a row for the chart, then the title, then the table
    assert interactive.anchor._from.row == 0
    assert at("issue-9", ws, "A2").value == "Region 1"
    assert interactive.series[0].val.numRef.f == f"'Interactive Chart'!$B${2 + p.row}:$E${2 + p.row}"

    sheets = {s["name"]: s for s in nt.preview(fixture("issue-9"))["sheets"]}
    assert [c["kind"] for c in sheets["Column and Bar Charts"]["charts"]] == ["column", "bar"]
    assert sheets["Interactive Chart"]["rowHeights"][0] > 100  # the row making room for the chart


def test_chart_links():
    # This chart's data was stored as a range of row IDs plus a reference to cells since deleted,
    # which Numbers ignores.
    result, wb = convert("issue-66-collab")
    assert result["charts"] == 1 and result["unsupported"]["charts"] == 0
    chart, = chart_on(wb["Credit"])
    assert chart.grouping == "stacked"
    assert [s.val.numRef.f for s in chart.series] == ["Credit!$B$10:$B$15", "Credit!$C$10:$C$15", "Credit!$D$10:$D$15"]

    # Cells in a run become a range; scattered cells on one sheet a union of ranges, as in Apple's
    # export; cells on several sheets, or a chart's own copy of values, go on a hidden sheet.
    columns = []

    class Data:
        def column(self, values, formats=None):
            columns.append(values)
            return ["Chart Data", 0, len(columns) - 1, len(values) - 1, len(columns) - 1]

    place = lambda cell: (cell[0], cell[1], cell[2])  # noqa: E731  (sheet, row, col) stand in for table cells
    info = lambda cell: (7, None)  # noqa: E731
    assert charts.reference([("S", 1, 1), ("S", 2, 1), ("S", 3, 1)], [], Data(), place, info) == ["S", 1, 1, 3, 1]
    assert charts.reference([("S", 1, 1), ("S", 1, 3), ("S", 1, 5)], [], Data(), place, info) == "=(S!$B$2,S!$D$2,S!$F$2)"
    assert charts.reference([("S", 1, 1), ("T", 1, 1)], [], Data(), place, info) == ["Chart Data", 0, 0, 1, 0]
    assert columns[0] == [("=S!$B$2", 7), ("=T!$B$2", 7)]
    assert charts.reference(None, [1.5, 2.5], Data(), place, info) == ["Chart Data", 0, 1, 1, 1]


def test_layout_arrangement():
    # Laying out tables as they sit on a Numbers sheet, along one axis: tables one below the other
    # get their own rows, with a row the size of the space between; tables beside each other share
    # rows, lining up with the nearest row edge.
    starts, gaps = layout.arrange([(0, [20] * 5), (178, [20] * 4)], 0, 409)
    assert starts == [0, 6] and gaps == {5: 78}
    starts, gaps = layout.arrange([(0, [20] * 5), (0, [20] * 3)], 0, 409)
    assert starts == [0, 0] and gaps == {}
    starts, _ = layout.arrange([(0, [20] * 5), (41, [20] * 3)], 0, 409)
    assert starts == [0, 2]
    # Tables sharing columns can't share rows: the second moves below the first, after a gap row.
    starts, gaps = layout.arrange([(0, [20] * 5), (50, [20] * 3)], 0, 409, lambda i, j: True)
    assert starts == [0, 6] and gaps == {5: layout.GAP}
    # A big space is split into rows no taller than Excel allows.
    starts, gaps = layout.arrange([(500, [20])], 0, 409)
    assert starts == [2] and gaps == {0: 250, 1: 250}


def test_one_sheet_layout():
    # Laid out as in Numbers, a Numbers sheet's tables share one worksheet, each where it sits on
    # the Numbers sheet, and formulas reading another table point at it there.
    options = {"default": "sheet"}
    result, wb = convert("create-formulas", options)
    assert result["sheets"] == 2 and wb.sheetnames == ["Main Sheet", "Powers Sheet"]
    ws = wb["Main Sheet"]
    plan = nt.load(fixture("create-formulas")).layout(options).sheets[0]
    tables = {p.table.name: p for p in plan.tables}
    tests, food, animals = tables["Formula Tests"], tables["Food Table"], tables["Animal Table"]
    assert food.row == tests.row and food.col > tests.col  # side by side, sharing rows
    assert animals.col == food.col and animals.row > food.row  # one below the other, sharing columns
    formula = at("create-formulas", ws, "A7", options, "Formula Tests").value
    assert formula == "=SUM(" + shifted("create-formulas", ws.title, "A2:B2", options, "Formula Tests") + ")"
    textjoin = next(c.value for row in ws.iter_rows() for c in row if str(c.value).startswith("=_xlfn.TEXTJOIN("))
    assert textjoin == '=_xlfn.TEXTJOIN(",",FALSE,' + nt.shift_rows("B1:D1", animals.row, animals.col) + ")"
    powers = next(c.value for row in ws.iter_rows() for c in row if str(c.value).startswith("=MEDIAN("))
    assert "'Powers Sheet'!" in powers  # a table on another Numbers sheet is on another worksheet

    sheets = nt.preview(fixture("create-formulas"), layout_options=options)["sheets"]
    assert [(s["name"], s["tableCount"], s["layout"]) for s in sheets] == [("Main Sheet", 9, "sheet"), ("Powers Sheet", 5, "sheet")]
    # Each Numbers sheet can be laid out its own way.
    mixed = {"default": "sheet", "sheets": {"Powers Sheet": "tables"}}
    names = [s["name"] for s in nt.preview(fixture("create-formulas"), layout_options=mixed)["sheets"]]
    assert names[0] == "Main Sheet" and len(names) == 6 and all(n.startswith("Powers Sheet - ") for n in names[1:])


def test_font_names():
    assert nt.font_name("Helvetica Neue") == "Arial" and nt.font_name("HelveticaNeue-Medium") == "Arial"
    assert nt.font_name("AvenirNext-DemiBold") == "Segoe UI" and nt.font_is_bold("AvenirNext-DemiBold")
    assert nt.font_name("Georgia-Bold") == "Georgia" and not nt.font_is_bold("Georgia")


def test_custom_formats_convert():
    convert("test-custom-formats")


def test_password_protected():
    for password, expected in ((None, "password-protected"), ("wrong", "isn't right")):
        try:
            nt.load(fixture("encrypted"), password)
        except nt.UserError as err:
            assert expected in str(err) and err.extra["needsPassword"]
            assert err.extra["hint"] == "s3cr3t"  # shown from the first prompt; this sample's hint is its password
        else:
            raise AssertionError("expected a UserError")
    assert nt.preview(fixture("encrypted"), "s3cr3t")["sheets"]
    convert_with = os.path.join(OUT, "encrypted.xlsx")
    nt.convert(fixture("encrypted"), convert_with, "s3cr3t")
    assert load_workbook(convert_with).worksheets


def test_folder_package():
    # Numbers on a Mac can save a file as a folder; those open directly.
    assert os.path.isdir(fixture("folder-package"))
    result, _ = convert("folder-package")
    assert result["sheets"] >= 1


def test_form_controls():
    _, wb = convert("test-actions")
    ws = wb.worksheets[0]
    rules = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    rule = lambda ref: rules[shifted("test-actions", ws.title, ref)]  # noqa: E731
    assert rule("A5").type == "list" and rule("A5").formula1 == '"Item 1,Item 2,Item 3"'  # pop-up menu
    assert rule("B2").type == "whole" and rule("B2").formula2 == "5"  # star rating 0-5
    assert rule("A3").type == "whole" and rule("A3").formula2 == "100"  # slider
    assert at("test-actions", ws, "A1").value is False and at("test-actions", ws, "B1").value is True  # checkboxes


def test_comments_links_bullets_and_rich_text():
    _, wb = convert("test-comments")
    note = at("test-comments", wb.worksheets[0], "A1").comment
    assert note.text.startswith("Comment #1") and note.author == "Jon Connell"

    _, wb = convert("test-hlinks")
    ws = wb.worksheets[0]
    assert at("test-hlinks", ws, "A1").hyperlink.target == "http://news.bbc.co.uk/"
    assert "http://google.co.uk/" in at("test-hlinks", ws, "A1").comment.text  # one link per cell; the rest are noted
    assert at("test-hlinks", ws, "A3").value == "this is a list:\n• first item\n• second item\n• third item"

    _, wb = convert("test-bullets")
    ws = wb.worksheets[0]
    assert at("test-bullets", ws, "B1").value == "• this\n• is some\n• bulleted\n• text"
    # B2 is read by =LEFT(B2,1); bullets are display-only in Numbers, so its text stays as it was.
    assert at("test-bullets", ws, "B2").value == "this\nis some\nnumbered\ntext"

    out = os.path.join(OUT, "rich.xlsx")
    nt.convert(fixture("test-formats"), out)
    ws = load_workbook(out, rich_text=True).worksheets[0]
    rich = at("test-formats", ws, "C7").value
    runs = [(str(part.text), bool(part.font.b), bool(part.font.i)) for part in rich if hasattr(part, "font")]
    assert ("the", True, False) in runs and ("quick", False, True) in runs

    preview = nt.preview(fixture("test-formats"))["sheets"][0]
    text, _, info = grid_cell(preview, placed("test-formats", preview["name"]), 6, 2)
    assert text == "the quick brown fox jumped" and len(info["runs"]) == 9


def test_hidden_rows_and_columns():
    # No sample file has hidden rows, so mark some hidden in the loaded file (as Numbers stores them).
    path = fixture("test-bgcolour")
    book = nt.load(path)
    table = book.tables[0][1]
    model = table._model
    bds = model.objects[table._table_id].base_data_store
    for bucket, index in ((bds.rowHeaders.buckets[0].identifier, 1), (bds.columnHeaders.identifier, 2)):
        headers = model.objects[bucket].headers
        header = next((h for h in headers if h.index == index), None) or headers.add()
        header.index, header.hidingState = index, 1
    book.layouts.clear()  # laid out before the change
    sheet = nt.preview(path)["sheets"][0]
    p = book.layout().place[table._table_id]
    assert sheet["hiddenRows"] == [p.row + 1] and sheet["hiddenCols"] == [p.col + 2]
    out = os.path.join(OUT, "hidden.xlsx")
    nt.convert(path, out)
    ws = load_workbook(out).worksheets[0]
    assert ws.row_dimensions[p.row + 2].hidden and ws.column_dimensions["C"].hidden
    nt._cache.clear()


def test_formulas_say_why_they_keep_values():
    def why(formula):
        reasons = []
        assert nt.translate_formula(formula, FakeLayout, "Main", why=reasons) is None
        return reasons[0]
    assert why('STOCK("AAPL",1)') == "function:STOCK"
    assert why("IF(A1)") == "arguments:IF"
    assert why("GCD(5,(A1,A2))") == "list"
    assert why('DATEDIF("15/05/2001","15/07/2002","D")') == "date-text"
    assert why('AVERAGEIFS({1,2},{1,0},">0")') == "array"
    assert why("SUM(Unknown::A1)") == "missing"
    assert why("SUM(#REF!)") == "broken"


def test_change_report():
    def report(name):
        return nt.preview(fixture(name))["report"]

    formulas = report("create-formulas")
    (kept,) = [c for c in formulas["changes"] if c["kind"] == "formulas"]
    assert kept["title"] == "1 formula keeps its last value" and "#REF!" in kept["detail"]
    cell = kept["cells"][0]
    sheet = next(s for s in nt.preview(fixture("create-formulas"))["sheets"] if s["name"] == cell["sheet"])
    shown = sheet["rows"][cell["row"]][cell["col"]]  # plain text, or [text, style, ...]
    assert (shown if isinstance(shown, str) else shown[0]) == "#REF!" and cell["ref"] == f"A{cell['row'] + 1}"
    assert "109 formulas that keep calculating in Excel" in formulas["kept"]

    charts_file = report("issue-9")
    kinds = {c["kind"]: c for c in charts_file["changes"]}
    assert list(kinds) == ["charts"] and "Bubble charts" in kinds["charts"]["detail"]
    assert "7 text boxes" in charts_file["kept"]  # converted since v1.9
    assert any(n["kind"] == "interactive" for n in charts_file["notes"])

    highlighted = report("issue-66-collab")
    assert not highlighted["changes"] and "Conditional highlighting in 1 table" in highlighted["kept"]
    custom = {c["kind"]: c for c in report("test-custom-formats")["changes"]}
    assert "DURATION" in custom["formulas"]["detail"] and custom["formats"]["title"].endswith("custom number format")
    pivot_report = report("test-pivot")
    assert not pivot_report["changes"] and not any(n["kind"] == "pivot-values" for n in pivot_report["notes"])
    assert "2 pivot tables as Excel PivotTables, with the same fields, groups, totals and order" in pivot_report["kept"]
    fonts = next(n for n in report("test-styles")["notes"] if n["kind"] == "fonts")
    assert "Helvetica Neue → Arial" in fonts["detail"]
    assert not report("test-styles")["changes"]


def test_unreadable_files_show_their_picture():
    bad = os.path.join(OUT, "Damaged.numbers")
    with zipfile.ZipFile(fixture("test-styles")) as zin, zipfile.ZipFile(bad, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, b"\x00broken" * 40 if item.filename.endswith("Document.iwa") else zin.read(item.filename))
    try:
        nt.load(bad)
    except nt.UserError as err:
        assert err.extra["reason"] == "unreadable" and err.extra["picture"].startswith("data:image/jpeg;base64,")
        assert "Damaged" not in err.extra["details"] and OUT not in err.extra["details"]
    else:
        raise AssertionError("expected a UserError")
    assert nt.technical(OSError(f"can't open {bad}"), bad) == "OSError: can't open <file>"


def test_contents_sheet():
    out = os.path.join(OUT, "contents.xlsx")
    result = nt.convert(fixture("create-formulas"), out, layout_options={"default": "sheet"}, contents=True)
    wb = load_workbook(out)
    assert result["contents"] and wb.sheetnames == ["Contents", "Main Sheet", "Powers Sheet"]
    ws = wb["Contents"]
    links = [row[2] for row in ws.iter_rows(min_row=5) if row[2].hyperlink]
    assert len(links) == 14 and all(link.hyperlink.location.startswith("'") for link in links)
    target = links[0].hyperlink.location  # e.g. 'Main Sheet'!H42: the table's top-left cell
    sheet, cell = target.split("!")
    assert wb[sheet.strip("'")][cell].value is not None
    plain = nt.convert(fixture("create-formulas"), os.path.join(OUT, "no-contents.xlsx"))
    assert not plain["contents"] and "Contents" not in load_workbook(os.path.join(OUT, "no-contents.xlsx")).sheetnames


def drawings_of(path):
    """{worksheet name: the XML of its drawing (charts, text boxes, shapes and pictures)} in a workbook."""
    found = {}
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        targets = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', z.read("xl/_rels/workbook.xml.rels").decode()))
        for name, rid in re.findall(r'<sheet name="([^"]+)" sheetId="\d+" r:id="(rId\d+)"', z.read("xl/workbook.xml").decode()):
            rels = "xl/" + targets[rid].replace("worksheets/", "worksheets/_rels/") + ".rels"
            drawing = re.search(r'Target="../drawings/(drawing\d+\.xml)"', z.read(rels).decode()) if rels in names else None
            if drawing:
                found[html.unescape(name)] = z.read("xl/drawings/" + drawing.group(1)).decode()
    return found


def test_text_boxes_convert():
    # issue-9 has a text box describing the charts on each sheet, with some words in bold: Excel gets
    # the same text, runs, font and padding, in a box that grows with its text as in Numbers.
    result, _ = convert("issue-9")
    assert result["drawings"] == 7 and result["unsupported"]["shapes"] == 0
    xml = drawings_of(os.path.join(OUT, "issue-9.xlsx"))["Column and Bar Charts"]
    (box,) = re.findall(r"<xdr:sp .*?</xdr:sp>", xml, re.S)
    assert 'txBox="1"' in box and "<a:spAutoFit/>" in box and 'lIns="50800"' in box  # Numbers' 4 pt padding
    runs = re.findall(r'<a:r><a:rPr [^>]*b="(\d)"[^>]*>.*?<a:t>([^<]*)</a:t>', box, re.S)
    assert runs[:3] == [("1", "Column "), ("0", "and "), ("1", "bar")]
    assert 'sz="1400"' in box and 'typeface="Arial"' in box  # 14 pt Helvetica Neue, as Windows has it
    assert len(re.findall(r"<xdr:graphicFrame", xml)) == 2  # the charts are still there
    # The table moved to column A, so the text box (left of it in Numbers) goes right of it, as a chart would.
    ws = next(ws for ws in nt.load(fixture("issue-9")).layout().sheets if ws.name == "Column and Bar Charts")
    (d, (x, y)), = ws.drawings
    assert d.kind == "text" and x > ws.rect(ws.tables[0])[2] * layout.PT_TO_PX and y == 30

    sheets = {s["name"]: s for s in nt.preview(fixture("issue-9"))["sheets"]}
    (shown,) = sheets["Column and Bar Charts"]["drawings"]
    assert shown["kind"] == "text" and shown["anchor"] == "t" and (shown["x"], shown["y"]) == (x, y)
    text, style = shown["paragraphs"][0]["runs"][0]
    assert text == "Column " and style["b"] and style["font"] == "Arial" and style["size"] == 14

    # Laid out as in Numbers, a table whose text box sat left of it keeps it there.
    options = {"default": "sheet"}
    ws = next(ws for ws in nt.load(fixture("issue-9")).layout(options).sheets if ws.name == "Bubble Chart")
    assert ws.tables[0].col == 1 and ws.drawings[0][1][0] == 0


def test_room_for_drawings():
    # A text box or picture above (or left of) everything else gets room, as charts do; one just a
    # little above doesn't move anything.
    table = layout.Footprint(None, (0.0, 130.0), [(20, False)], [(50, False)], None)
    title = SimpleNamespace(frame=(0.0, 0.0, 540.0, 40.0))
    assert layout.corner([table], [], [title]) == (0.0, 0.0)
    near = SimpleNamespace(frame=(0.0, 126.0, 100.0, 20.0))
    assert layout.corner([table], [], [near]) == (0.0, 130.0)
    chart = SimpleNamespace(frame=(10.0, 60.0, 100.0, 50.0))
    assert layout.corner([table], [chart], []) == (0.0, 54.0)
    assert layout.corner([], [], [title]) == (0.0, 0.0)  # a sheet with only a text box


def test_images_convert():
    # images.numbers (made by tools/make_images_fixture.py, as no sample from Numbers has pictures)
    # has a picture with a border, one cropped by a rounded mask, one saved as TIFF (Excel gets the
    # preview Numbers keeps) and one as HEIC (no preview, so it's reported), a group holding a
    # picture and an oval, and a rounded rectangle turned 30 degrees.
    result, _ = convert("images")
    assert result["drawings"] == 6 and result["unsupported"] == {"charts": 0, "images": 1, "shapes": 0, "pivots": 0}
    path = os.path.join(OUT, "images.xlsx")
    xml = drawings_of(path)["Sheet 1"]
    pictures = re.findall(r"<xdr:pic>.*?</xdr:pic>", xml, re.S)
    assert len(pictures) == 4
    assert '<a:ln w="25400">' in pictures[0]  # the 2 pt border
    assert 'descr="A sunrise"' in pictures[0] and "<a:hlinkClick" in pictures[0]  # its description and link
    assert '<a:srcRect l="25000" t="25000" r="25000" b="25000"/>' in pictures[1] and "<a:custGeom>" in pictures[1]
    shapes = re.findall(r"<xdr:sp .*?</xdr:sp>", xml, re.S)
    turned = next(s for s in shapes if "Turned" in s)
    assert 'rot="19800000"' in turned  # 30 degrees anticlockwise in Numbers is 330 clockwise in Excel
    oval = next(s for s in shapes if "Grouped" in s)
    assert oval.count("<a:cubicBezTo>") == 4 and "<a:solidFill>" in oval
    with zipfile.ZipFile(path) as z:
        assert len([n for n in z.namelist() if n.startswith("xl/media/")]) == 4

    ws = nt.load(fixture("images")).layout().sheets[0]
    placed = {d.name or d.image_name: position for d, position in ws.drawings}
    assert placed["Grouped.png"] == (0, 200) and placed["Grouped"] == (107, 200)  # where their group is
    preview = nt.preview(fixture("images"))
    assert [c["kind"] for c in preview["report"]["changes"]] == ["images"]
    assert "2 shapes and 4 images" in preview["report"]["kept"]
    pictures = [d for d in preview["sheets"][0]["drawings"] if d["kind"] == "image"]
    assert all(d["src"].startswith("data:image/png;base64,") for d in pictures) and pictures[1]["crop"] == [0.25] * 4
    budget = [10]  # too little for any picture: they're drawn as placeholders
    assert "src" not in drawings.describe(ws.drawings[0][0], (0, 0), budget)


def test_shape_outlines():
    # Numbers stretches a bezier outline to the shape's natural size, whatever size it was drawn at.
    d = drawings.Drawing(1, "shape")
    line = TSD.PathSourceArchive(bezier_path_source={"naturalSize": {"width": 540, "height": 0}, "path": {"elements": [
        {"type": 1, "points": [{"x": 0, "y": 0}]}, {"type": 2, "points": [{"x": 501, "y": 0}]}]}})
    assert drawings.read_outline(d, line) == (540, 0) and d.outline == "line"
    assert drawings.line_frame(d, (10.0, 20.0, 540.0, 0.0)) == (10.0, 20.0, 540.0, 0.0)
    box = TSD.PathSourceArchive(bezier_path_source={"naturalSize": {"width": 300, "height": 50}, "path": {"elements": [
        {"type": 1, "points": [{"x": 0, "y": 0}]}, {"type": 2, "points": [{"x": 100, "y": 0}]},
        {"type": 2, "points": [{"x": 100, "y": 100}]}, {"type": 2, "points": [{"x": 0, "y": 100}]},
        {"type": 5}, {"type": 1, "points": [{"x": 0, "y": 0}]}]}})
    drawings.read_outline(d, box)
    assert d.outline == "rect"
    star = TSD.PointPathSourceArchive(type=100, point={"x": 5, "y": 0.5}, naturalSize={"width": 100, "height": 100})
    path, size = drawings.point_path(star)
    assert size == (100, 100) and [cmd for cmd, _ in path] == ["M"] + ["L"] * 9 + ["Z"]
    assert [round(v, 6) for v in path[0][1][0]] == [50, 0]  # the top point
    rounded = TSD.ScalarPathSourceArchive(type=0, scalar=10, naturalSize={"width": 100, "height": 40})
    assert [cmd for cmd, _ in drawings.scalar_path(rounded)[0]].count("C") == 4
    bubble = TSD.CalloutPathSourceArchive(natural_size={"width": 100, "height": 40}, tail_position={"x": 20, "y": 70},
                                          tail_size=10, corner_radius=5)
    path, _ = drawings.callout_path(bubble)
    assert ("L", [(20, 70)]) in path  # the tail's tip, below the box


def test_highlighting_converts():
    # issue-66-collab colours negative numbers red: Excel gets the same rule on the same cells.
    _, wb = convert("issue-66-collab")
    areas = lambda cf: " ".join(sorted(str(r) for r in cf.sqref.ranges))  # noqa: E731
    (rules,) = wb["Credit"].conditional_formatting
    assert areas(rules) == "B10:E15 E9"
    (rule,) = rules.rules
    assert (rule.type, rule.operator, rule.formula, rule.stopIfTrue) == ("cellIs", "lessThan", ["0"], None)
    assert rule.dxf.font.color.rgb == "FFFF0000"

    # Other kinds of rule, in the order Numbers tries them; Excel stops at the first that matches, as
    # Numbers does. One Excel can't calculate is left out and reported.
    book = nt.load(fixture("issue-66-collab"))
    found = next(iter(book.highlights.values()))
    key = next(iter(found.sets))
    cell = found.sets[key][0]["nodes"][0]  # LINKED_CELL_REF_NODE: the cell being tested
    found.sets[key] = [{"nodes": [cell] + rpn("Done", "="), "style": {"b": 1}},
                       {"nodes": [cell] + rpn("ISBLANK/1"), "style": {"bg": "#FFFF00"}},
                       {"nodes": [cell] + rpn(600, ">"), "style": {"color": "#0000FF"}},
                       {"nodes": [cell] + rpn("STOCK/1"), "style": {"i": 1}}]
    book.layouts.clear()
    try:
        out = os.path.join(OUT, "highlights.xlsx")
        nt.convert(fixture("issue-66-collab"), out)
        rules = sorted((rule.priority, areas(cf), rule.type, rule.operator, rule.formula, rule.stopIfTrue)
                       for cf in load_workbook(out)["Credit"].conditional_formatting for rule in cf.rules)
        assert rules == [(1, "B10:E15 E9", "cellIs", "equal", ['"Done"'], True),
                         (2, "E9", "expression", None, ["ISBLANK(E9)"], True),
                         (3, "B10:E15", "expression", None, ["ISBLANK(B10)"], True),
                         (4, "B10:E15 E9", "cellIs", "greaterThan", ["600.0"], True)]
        report = nt.preview(fixture("issue-66-collab"))
        assert [c["kind"] for c in report["report"]["changes"]] == ["highlighting"]
        # The preview shows what "Cell Value" rules highlight: B10 holds 629.
        sheet = report["sheets"][0]
        _, style, *_ = sheet["rows"][9][1]
        assert sheet["styles"][style]["color"] == "#0000FF"
    finally:
        nt._cache.clear()
    assert highlights.matches("<", None, "0") is False and highlights.matches("<=", None, "0")  # a blank cell is 0
    assert highlights.matches("==", "done", '"Done"') and highlights.matches(">", "text", "100")
    assert highlights.blocks([(8, 4)] + [(r, c) for r in range(9, 15) for c in range(1, 5)]) == [(8, 4, 8, 4), (9, 1, 14, 4)]


def test_categories_convert():
    # Tables grouped into categories come out as Numbers shows them: a category column, a row with
    # each group's name, its rows under it in Numbers' order, groups inside groups, and each group
    # an Excel outline group that collapses (its button on the group's row, above its rows).
    result, wb = convert("test-categories")
    ws = wb["Categories - Categories"]
    assert [ws.cell(r, 1).value for r in (3, 14, 25)] == ["Transport", "Fruit", "Animal"]
    assert [ws.cell(2, c).value for c in range(1, 5)] == [None, "Description", "Category", "Count"]
    # Transport's rows are the table's rows 1, 5, 6...: each still reads its own row of the table
    # it copies from (whose title row puts row 1 on row 3).
    assert ws["B4"].value == "='Categories - Uncategorized'!A3" and ws["B5"].value == "='Categories - Uncategorized'!A7"
    assert ws.row_dimensions[3].outlineLevel == 0 and ws.row_dimensions[4].outlineLevel == 1
    assert ws.sheet_properties.outlinePr.summaryBelow is False
    assert result["formulasAsValues"] == 0

    nested = wb["Categories - Nested Categories"]
    assert [nested.cell(r, 1).value for r in (3, 4)] == ["Not “Edible”", "Transport"]  # a checkbox's groups
    assert nested["A4"].alignment.indent == 1 and nested.row_dimensions[5].outlineLevel == 2
    assert wb["Categories - Date Year-Month"]["A3"].value == "February 2010"  # by a part of a date
    assert wb["Categories - Date Year-Quarter"]["A3"].value == "2010-Q1"
    assert wb["Categories - Date Weekday"]["A3"].value == "Sunday"

    # On one sheet the tables sit side by side, so their rows can't collapse; they're still grouped.
    _, wb = convert("test-categories", {"default": "sheet"})
    p = placed("test-categories", "Categories", {"default": "sheet"}, table="Categories")
    ws = wb["Categories"]
    assert ws.cell(p.row + 2, p.col + 1).value == "Transport"
    assert not any(ws.row_dimensions[r].outlineLevel for r in range(1, 40))

    preview = nt.preview(fixture("test-categories"))
    sheet = next(s for s in preview["sheets"] if s["name"] == "Categories - Categories")
    assert sheet["rows"][2][0][0] == "Transport" and sheet["rows"][3][1][0] == "Airplane"
    assert "10 tables grouped into categories, with each group’s row and summaries" in preview["report"]["kept"]
    assert not [c for c in preview["report"]["changes"] if c["kind"].startswith("categor")]


def test_category_summaries():
    # The sample's groups have no summaries, so give the Categories table's groups a sum of Count
    # (Numbers' code for Sum): Excel gets a label row, then each group's sum as a live formula
    # over its rows.
    path = fixture("test-categories")
    doc = Document(path)
    table = next(t for s in doc.sheets for t in s.tables if t.name == "Categories")
    model = doc._model
    tm = model.objects[table._table_id]
    archive = model.objects[model.objects[tm.category_owner.identifier].group_by[0].identifier]
    uids = model.objects[tm.base_column_row_uids.identifier]
    summary = archive.column_agg_type.add()
    summary.column_uid.CopyFrom(uids.sorted_column_uids[uids.column_uid_for_index[2]])
    summary.level, summary.agg_type = 1, 2
    nt._cache[(path, os.path.getmtime(path), None)] = nt.Book(doc)
    try:
        _, wb = convert("test-categories")
        ws = wb["Categories - Categories"]
        assert (ws["A3"].value, ws["D3"].value) == ("Category:", "Sum:")
        assert [(ws.cell(r, 1).value, ws.cell(r, 4).value) for r in (4, 15, 26)] == [
            ("Transport", "=SUM(D5:D14)"), ("Fruit", "=SUM(D16:D25)"), ("Animal", "=SUM(D27:D36)")]
        assert ws["A4"].font.b and ws["D4"].font.b
    finally:
        nt._cache.clear()

    # A summary whose code isn't known is recognised from the results Numbers saved, and kept as
    # they are when they don't tell one function from another, or disagree with the known code.
    groups = [SimpleNamespace(uid="a", rows=[0, 1]), SimpleNamespace(uid="b", rows=[2])]
    data = [[2.0], [4.0], [9.0]]
    tools = SimpleNamespace(number=lambda value: value, filled=lambda value: value is not None)
    saved = {("a", 0): (3.0, None), ("b", 0): (9.0, None)}
    assert categories.settle(categories.UNKNOWN, groups, 0, saved, data, tools) == "average"
    assert categories.settle("sum", groups, 0, saved, data, tools) is None
    assert categories.settle(categories.UNKNOWN, groups[1:], 0, saved, data, tools) is None
    counts = {("a", 0): (2.0, None), ("b", 0): (1.0, None)}
    assert categories.settle(categories.UNKNOWN, groups, 0, counts, data, tools) == "count"
    assert categories.settle("sum", groups, 0, {}, data, tools) == "sum"  # nothing saved to check against


def test_references_into_grouped_tables():
    # Formulas reading a grouped table follow its rows to where they went: a range whose rows come
    # apart becomes a union of the pieces in functions that read one, and keeps its value elsewhere.
    book = nt.load(fixture("test-categories"))
    layout_ = book.layout()
    table = next(t for _, t in book.tables if t.name == "Categories")
    tid, p = table._table_id, layout_.place[table._table_id]
    other = f"'{p.sheet}'!"
    r, c = p.cell(5, 2)
    assert nt.translate_formula(f"⟦{tid}|5|5|2|2|0⟧+1", layout_, "Categories") == f"={other}{nt.column_letters(c)}{r + 1}+1"
    assert nt.translate_formula(f"SUM(⟦{tid}|5|7|2|2|0⟧)", layout_, "Categories") == f"=SUM({other}D5:D7)"
    assert nt.translate_formula(f"SUM(⟦{tid}|||2|2|0⟧)", layout_, "Categories") == (
        f"=SUM(({other}D4:D13,{other}D15:D24,{other}D26:D35))")
    why = []
    assert nt.translate_formula(f"VLOOKUP(1, ⟦{tid}|1|30|1|2|0⟧, 2, 0)", layout_, "Categories", why=why) is None
    assert why == ["category"]
    why = []
    assert nt.translate_formula(f"ROW(⟦{tid}|5|5|2|2|0⟧)", layout_, "Categories", why=why) is None and why == ["category"]


if __name__ == "__main__":
    tests = [(name, fn) for name, fn in globals().items() if name.startswith("test_")]
    for name, fn in tests:
        fn()
        print("ok", name)
    print(f"All {len(tests)} helper tests passed.")
